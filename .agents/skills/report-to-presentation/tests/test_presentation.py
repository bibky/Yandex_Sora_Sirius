
"""Unit tests for PowerPoint presentation generation."""

import sys
from pathlib import Path

import pytest
from pptx import Presentation


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from generate_presentation import generate_presentation, generate_from_json


@pytest.fixture
def sample_files(tmp_path):
    """Create temporary log and report files."""
    log_file = tmp_path / "application.log"
    log_file.write_text(
        "2026-10-10 INFO Application started\n"
        "2026-10-10 ERROR Database connection failed\n"
        "2026-10-10 WARNING High memory usage\n",
        encoding="utf-8",
    )

    report_file = tmp_path / "report.md"
    report_file.write_text(
        "# Weekly Report\n"
        "## Executive Summary\n"
        "The application experienced database issues.\n"
        "## Key Metrics\n"
        "- Total requests: 1500\n"
        "## Issues\n"
        "- Database connection failures\n"
        "## Recommendations\n"
        "- Improve database monitoring\n",
        encoding="utf-8",
    )

    return log_file, report_file


def test_presentation_file_created(sample_files, tmp_path):
    """Verify that a PowerPoint file is created."""
    log_file, report_file = sample_files
    output_file = tmp_path / "presentation.pptx"

    result = generate_presentation(
        log_file, report_file, output_file
    )

    assert result.exists()
    assert result.suffix == ".pptx"


def test_presentation_slide_count(sample_files, tmp_path):
    """Verify that the presentation contains six slides."""
    log_file, report_file = sample_files
    output_file = tmp_path / "presentation.pptx"

    generate_presentation(log_file, report_file, output_file)

    presentation = Presentation(output_file)

    assert len(presentation.slides) == 6


def test_presentation_slide_titles(sample_files, tmp_path):
    """Verify that the presentation has expected slide titles."""
    log_file, report_file = sample_files
    output_file = tmp_path / "presentation.pptx"

    generate_presentation(log_file, report_file, output_file)

    presentation = Presentation(output_file)
    titles = [
        slide.shapes.title.text
        for slide in presentation.slides
    ]

    assert titles == [
        "Application Performance Report",
        "Executive Summary",
        "Key Metrics",
        "Errors and Critical Events",
        "Main Findings",
        "Recommendations",
    ]


def test_presentation_contains_log_error(sample_files, tmp_path):
    """Verify that log errors appear in the presentation."""
    log_file, report_file = sample_files
    output_file = tmp_path / "presentation.pptx"

    generate_presentation(log_file, report_file, output_file)

    presentation = Presentation(output_file)

    slide_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )

    assert "Database connection failed" in slide_text


def test_presentation_missing_input(tmp_path):
    """Verify that missing input files raise an error."""
    with pytest.raises(FileNotFoundError):
        generate_presentation(
            tmp_path / "missing.log",
            tmp_path / "missing.md",
            tmp_path / "presentation.pptx",
        )



def test_generate_from_json_creates_russian_presentation(tmp_path):
    """Verify that JSON input produces slides with Russian text."""
    import json
    from pptx import Presentation

    presentation_data = {
        "title": "Отчёт о работе приложения",
        "subtitle": "Тестовая презентация",
        "language": "ru",
        "slides": [
            {
                "title": "Краткое резюме",
                "bullets": [
                    "Приложение работает стабильно.",
                    "Обнаружено одно предупреждение.",
                ],
            }
        ],
    }

    json_path = tmp_path / "presentation.json"
    output_path = tmp_path / "presentation.pptx"

    json_path.write_text(
        json.dumps(presentation_data, ensure_ascii=False),
        encoding="utf-8",
    )

    result = generate_from_json(json_path, output_path)

    assert result == output_path
    assert output_path.exists()

    presentation = Presentation(output_path)

    assert len(presentation.slides) == 2

    all_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )

    assert "Отчёт о работе приложения" in all_text
    assert "Краткое резюме" in all_text
    assert "Приложение работает стабильно." in all_text



def test_generate_from_json_does_not_overwrite_existing_file(tmp_path):
    """Verify that the generator does not overwrite an existing file."""
    import json
    import pytest

    presentation_data = {
        "title": "Test Presentation",
        "slides": [
            {
                "title": "Summary",
                "bullets": ["Test content"],
            }
        ],
    }

    json_path = tmp_path / "presentation.json"
    output_path = tmp_path / "existing.pptx"

    json_path.write_text(
        json.dumps(presentation_data),
        encoding="utf-8",
    )

    output_path.write_bytes(b"Existing presentation content")

    with pytest.raises(FileExistsError):
        generate_from_json(json_path, output_path)

    assert output_path.read_bytes() == b"Existing presentation content"



def test_generate_from_json_rejects_invalid_structure(tmp_path):
    """Verify that invalid presentation data is rejected."""
    import json
    import pytest

    presentation_data = {
        "title": "Test Presentation",
        "slides": "This should be a list",
    }

    json_path = tmp_path / "invalid.json"
    output_path = tmp_path / "presentation.pptx"

    json_path.write_text(
        json.dumps(presentation_data),
        encoding="utf-8",
    )

    with pytest.raises(
        ValueError,
        match="Presentation slides must be a non-empty list",
    ):
        generate_from_json(json_path, output_path)

    assert not output_path.exists()

