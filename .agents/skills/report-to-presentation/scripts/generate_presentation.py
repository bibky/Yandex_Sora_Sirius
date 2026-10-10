
"""Generate PowerPoint presentations from logs and reports."""

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from parse_logs import parse_logs
from parse_reports import parse_reports


def add_title_slide(prs, title, subtitle):
    """Add a title slide to the presentation."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_content_slide(prs, title, lines):
    """Add a slide containing bullet-point text."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    for index, line in enumerate(lines):
        paragraph = (
            text_frame.paragraphs[0]
            if index == 0
            else text_frame.add_paragraph()
        )
        paragraph.text = str(line)
        paragraph.level = 0
        paragraph.font.size = Pt(18)


def find_report_section(report_data, title):
    """Find a report section by its title."""
    for section in report_data["sections"]:
        if section["title"].lower() == title.lower():
            return section["content"].splitlines()
    return []


def generate_presentation(log_path, report_path, output_path):
    """Build a PowerPoint presentation from logs and a report."""
    log_data = parse_logs(log_path)
    report_data = parse_reports(report_path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Application Performance Report",
        "Generated from application logs and text reports",
    )

    summary = find_report_section(report_data, "Executive Summary")
    add_content_slide(
        prs,
        "Executive Summary",
        summary or ["No executive summary available."],
    )

    metrics = [
        f"Total log lines: {log_data['total_lines']}",
        f"Critical events: {log_data['counts']['CRITICAL']}",
        f"Errors: {log_data['counts']['ERROR']}",
        f"Warnings: {log_data['counts']['WARNING']}",
        f"Info messages: {log_data['counts']['INFO']}",
    ]

    metrics.extend(find_report_section(report_data, "Key Metrics"))
    add_content_slide(prs, "Key Metrics", metrics)

    errors = [
        entry["message"]
        for entry in log_data["errors"]
    ]
    add_content_slide(
        prs,
        "Errors and Critical Events",
        errors or ["No errors or critical events detected."],
    )

    findings = find_report_section(report_data, "Issues")
    add_content_slide(
        prs,
        "Main Findings",
        findings or ["No issues documented in the report."],
    )

    recommendations = find_report_section(
        report_data, "Recommendations"
    )
    add_content_slide(
        prs,
        "Recommendations",
        recommendations or ["No recommendations available."],
    )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)

    return output



def generate_from_json(json_path, output_path):
    """Generate a PowerPoint presentation from structured JSON data."""
    source = Path(json_path)

    if not source.is_file():
        raise FileNotFoundError(f"JSON file not found: {source}")

    with source.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if not isinstance(data, dict):
        raise ValueError("Presentation data must be a JSON object.")

    title = data.get("title")
    slides = data.get("slides")

    if not isinstance(title, str) or not title.strip():
        raise ValueError("Presentation title must be a non-empty string.")

    if not isinstance(slides, list) or not slides:
        raise ValueError("Presentation slides must be a non-empty list.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        title,
        str(data.get("subtitle", "")),
    )

    for slide_data in slides:
        if not isinstance(slide_data, dict):
            raise ValueError("Each slide must be a JSON object.")

        slide_title = slide_data.get("title")
        bullets = slide_data.get("bullets")

        if not isinstance(slide_title, str) or not slide_title.strip():
            raise ValueError("Each slide must have a non-empty title.")

        if not isinstance(bullets, list) or not all(
            isinstance(bullet, str) for bullet in bullets
        ):
            raise ValueError("Slide bullets must be a list of strings.")

        add_content_slide(prs, slide_title, bullets)

    output = Path(output_path)

    if output.exists():
        raise FileExistsError(
            f"Output file already exists: {output}"
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)

    return output


def main():
    """Run the presentation generator from the command line."""
    parser = argparse.ArgumentParser(
        description="Generate a PowerPoint presentation."
    )

    parser.add_argument("--logs")
    parser.add_argument("--report")
    parser.add_argument("--input-json")
    parser.add_argument("--output", required=True)

    args = parser.parse_args()

    if args.input_json:
        if args.logs or args.report:
            parser.error(
                "--input-json cannot be combined with --logs or --report."
            )

        output = generate_from_json(
            args.input_json,
            args.output,
        )
    else:
        if not args.logs or not args.report:
            parser.error(
                "Provide --input-json or both --logs and --report."
            )

        output = generate_presentation(
            args.logs,
            args.report,
            args.output,
        )

    print(f"Presentation created: {output}")


if __name__ == "__main__":
    main()
